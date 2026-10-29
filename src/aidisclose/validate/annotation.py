"""Annotation files and inter-annotator agreement.

* :func:`write_annotation_file` writes a sample for annotators as an Excel
  workbook (drop-down lists for ``gold_state`` and ``gold_reason``, wrapped
  text, a codebook sheet) or as a UTF-8 CSV with byte-order mark, which Excel
  opens with the right encoding.
* :func:`read_annotation_file` reads either format back.
* :func:`to_label_studio` / :func:`from_label_studio` convert a sample to
  Label Studio tasks (with :data:`LABEL_STUDIO_CONFIG`) and a Label Studio JSON
  export back to AIDisclose annotations, one row per annotator.
* :func:`agreement_report` computes percent agreement, Cohen's kappa (two
  annotators), Fleiss' kappa and Krippendorff's alpha (nominal; any number of
  annotators, missing ratings allowed).
"""

from __future__ import annotations

import json
from collections import Counter
from collections.abc import Mapping, Sequence
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd

from ..models import REASONS_BY_STATE, STATES
from .metrics import cohen_kappa, confusion_matrix

REASONS = [r.value for state in REASONS_BY_STATE for r in REASONS_BY_STATE[state]]
_STATE_HELP = {
    "SUBSTANTIVE": "Focal firm uses/sells/invested in concrete AI AND a specific element (application, product, "
    "quantity, date, investment).",
    "WEAK": "Focal firm's own AI, but forward-looking, hedged, risk boilerplate or unspecific.",
    "NO_CLAIM": "No valid claim about the focal firm's own AI (third parties, markets, regulation, negation, "
    "false positives, headings).",
}


# ----------------------------------------------------------------- files
def write_annotation_file(df: pd.DataFrame, path: str | Path) -> Path:
    """Write an annotation sheet as ``.xlsx`` (with drop-downs) or ``.csv`` (UTF-8 with BOM)."""
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        _write_xlsx(df, path)
    else:
        df.to_csv(path, index=False, encoding="utf-8-sig")
    return path


def _write_xlsx(df: pd.DataFrame, path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.datavalidation import DataValidation
    except ImportError as exc:  # pragma: no cover - openpyxl is a dependency
        raise ImportError("writing .xlsx needs openpyxl (pip install openpyxl)") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "annotate"
    cols = list(df.columns)
    ws.append(cols)
    for row in df.itertuples(index=False):
        ws.append(["" if (isinstance(v, float) and v != v) or v is pd.NA else v for v in row])
    header_fill = PatternFill("solid", fgColor="DDDDDD")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    widths = {"text": 80, "prev_text": 45, "next_text": 45, "notes": 30, "firm_name": 22, "section": 14}
    n = len(df) + 1
    for i, c in enumerate(cols, 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = widths.get(c, 14)
        if c in ("text", "prev_text", "next_text", "notes"):
            for r in range(2, n + 1):
                ws[f"{letter}{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    lists = wb.create_sheet("lists")
    for i, s in enumerate(STATES, 1):
        lists.cell(row=i, column=1, value=s)
    for i, r in enumerate(REASONS, 1):
        lists.cell(row=i, column=2, value=r)
    lists.sheet_state = "hidden"
    for col, ref in (("gold_state", f"lists!$A$1:$A${len(STATES)}"), ("gold_reason", f"lists!$B$1:$B${len(REASONS)}")):
        if col in cols:
            letter = get_column_letter(cols.index(col) + 1)
            dv = DataValidation(type="list", formula1=f"={ref}", allow_blank=True, showErrorMessage=True)
            dv.error = "Choose a value from the list (see the codebook sheet)."
            ws.add_data_validation(dv)
            dv.add(f"{letter}2:{letter}{max(n, 2)}")
    ws.freeze_panes = "A2"
    cb = wb.create_sheet("codebook")
    cb.append(["state", "definition"])
    for s in STATES:
        cb.append([s, _STATE_HELP[s]])
    cb.append([])
    cb.append(["reason", "valid for state"])
    for state, reasons in REASONS_BY_STATE.items():
        for r in reasons:
            cb.append([r.value, state.value])
    cb.column_dimensions["A"].width = 20
    cb.column_dimensions["B"].width = 110
    wb.save(path)


def read_annotation_file(path: str | Path) -> pd.DataFrame:
    """Read an annotation sheet written by :func:`write_annotation_file` (or edited by hand)."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        df = pd.read_excel(path, sheet_name=0, dtype=str, keep_default_na=False)
    else:
        df = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]
    return df


# ---------------------------------------------------------- Label Studio
def _choices(name: str, values: Sequence[str], required: bool) -> str:
    req = ' required="true"' if required else ""
    inner = "".join(f'<Choice value="{escape(v)}"/>' for v in values)
    return f'  <Choices name="{name}" toName="target" choice="single"{req} showInline="true">{inner}</Choices>\n'


LABEL_STUDIO_CONFIG = (
    "<View>\n"
    '  <Header value="$firm_name ($section, FY $fiscal_year)"/>\n'
    '  <Text name="prev" value="$prev_text"/>\n'
    '  <Header value="Target sentence" size="5"/>\n'
    '  <Text name="target" value="$text"/>\n'
    '  <Text name="next" value="$next_text"/>\n'
    + _choices("state", STATES, True)
    + _choices("reason", REASONS, False)
    + '  <TextArea name="notes" toName="target" rows="2" placeholder="notes (optional)"/>\n'
    "</View>\n"
)


def to_label_studio(df: pd.DataFrame) -> list[dict[str, Any]]:
    """Convert an annotation sample to Label Studio tasks (import as JSON)."""
    fields = ["sentence_id", "firm_name", "fiscal_year", "section", "prev_text", "text", "next_text"]
    tasks = []
    for row in df.to_dict(orient="records"):
        data = {k: ("" if row.get(k) is None or pd.isna(row.get(k)) else str(row.get(k))) for k in fields}
        tasks.append({"data": data})
    return tasks


def from_label_studio(export: str | Path | bytes | list) -> pd.DataFrame:
    """Read a Label Studio JSON export into ``sentence_id, gold_state, gold_reason, annotator, notes``.

    One row per non-cancelled annotation, so tasks labelled by two annotators
    yield two rows (use :func:`agreement_report` on them).
    """
    if isinstance(export, bytes):
        data = json.loads(export.decode("utf-8"))
    elif isinstance(export, str) and export.lstrip().startswith(("[", "{")):
        data = json.loads(export)
    elif isinstance(export, (str, Path)):
        data = json.loads(Path(export).read_text(encoding="utf-8"))
    else:
        data = export
    if isinstance(data, dict):  # a single task
        data = [data]
    rows = []
    context_fields = ("firm_name", "fiscal_year", "section", "prev_text", "text", "next_text")
    available = [field for field in context_fields if any(field in (task.get("data") or {}) for task in data)]
    for task in data:
        task_data = task.get("data") or {}
        sid = str(task_data.get("sentence_id", task.get("id", "")))
        for ann in task.get("annotations") or []:
            if ann.get("was_cancelled"):
                continue
            who = ann.get("completed_by")
            if isinstance(who, Mapping):
                who = who.get("email") or who.get("id")
            row = {"sentence_id": sid, "gold_state": "", "gold_reason": "", "annotator": str(who or ""), "notes": ""}
            row.update({field: task_data.get(field, "") for field in available})
            for res in ann.get("result") or []:
                value = res.get("value") or {}
                name = res.get("from_name")
                if name in ("state", "reason") and value.get("choices"):
                    row["gold_" + name] = str(value["choices"][0])
                elif name == "notes" and value.get("text"):
                    row["notes"] = " ".join(value["text"])
            rows.append(row)
    return pd.DataFrame(rows, columns=["sentence_id", "gold_state", "gold_reason", "annotator", "notes", *available])


# -------------------------------------------------------------- agreement
def fleiss_kappa(ratings: pd.DataFrame) -> float:
    """Fleiss' kappa for items (rows) rated by the same number of annotators (columns)."""
    complete = ratings.dropna()
    if len(complete) < 2 or complete.shape[1] < 2:
        return float("nan")
    cats = sorted(set(complete.to_numpy().ravel()))
    n = complete.shape[1]
    counts = np.array([[list(r).count(c) for c in cats] for r in complete.itertuples(index=False)], dtype=float)
    p_i = (np.sum(counts**2, axis=1) - n) / (n * (n - 1))
    p_bar = p_i.mean()
    p_j = counts.sum(axis=0) / counts.sum()
    p_e = np.sum(p_j**2)
    return float((p_bar - p_e) / (1 - p_e)) if p_e < 1 else float("nan")


def krippendorff_alpha_nominal(ratings: pd.DataFrame) -> float:
    """Krippendorff's alpha for nominal data; rows are items, columns annotators, NaN = missing."""
    coincid: Counter = Counter()
    for row in ratings.itertuples(index=False):
        vals = [v for v in row if v is not None and v == v and str(v) != ""]
        m = len(vals)
        if m < 2:
            continue
        for i, a in enumerate(vals):
            for j, b in enumerate(vals):
                if i != j:
                    coincid[(a, b)] += 1.0 / (m - 1)
    if not coincid:
        return float("nan")
    cats = sorted({a for a, _ in coincid})
    n_c = {c: sum(v for (a, _), v in coincid.items() if a == c) for c in cats}
    n = sum(n_c.values())
    d_o = sum(v for (a, b), v in coincid.items() if a != b)
    d_e = sum(n_c[a] * n_c[b] for a in cats for b in cats if a != b) / (n - 1)
    return float(1 - d_o / d_e) if d_e > 0 else float("nan")


def agreement_report(
    annotations: Sequence[pd.DataFrame] | pd.DataFrame,
    on: str = "sentence_id",
    column: str = "gold_state",
    annotator_column: str = "annotator",
) -> dict[str, Any]:
    """Inter-annotator agreement on ``column``.

    Pass either a list of annotation tables (one per annotator) or one long
    table with an ``annotator`` column (e.g. from :func:`from_label_studio`).
    """
    if isinstance(annotations, pd.DataFrame):
        long = annotations.copy()
        if (
            annotator_column not in long
            or long[annotator_column].isna().any()
            or long[annotator_column].astype(str).str.strip().eq("").any()
        ):
            raise ValueError("each rating needs a nonempty annotator identifier")
        tables = {str(k): g for k, g in long.groupby(annotator_column)}
    else:
        tables = {f"annotator_{i + 1}": t for i, t in enumerate(annotations)}
    if len(tables) < 2:
        raise ValueError("agreement needs at least two annotators")
    wide = None
    for name, t in tables.items():
        s = t[[on, column]].copy()
        s[column] = s[column].fillna("").astype(str).str.strip().str.upper().str.replace(" ", "_")
        s = s[s[column] != ""]
        if s[on].isna().any() or s[on].astype(str).str.strip().eq("").any():
            raise ValueError(f"{name} has a rating without an item identifier")
        if s[on].duplicated().any():
            raise ValueError(f"{name} has duplicate ratings for an item")
        if column == "gold_state" and not s[column].isin(STATES).all():
            raise ValueError(f"{name} contains invalid gold_state labels")
        s = s.rename(columns={column: name})
        wide = s if wide is None else wide.merge(s, on=on, how="outer")
    wide = wide.set_index(on)
    names = list(tables)
    both = wide.dropna()
    rep: dict[str, Any] = {
        "annotators": names,
        "column": column,
        "n_items": int(len(wide)),
        "n_items_all_annotators": int(len(both)),
        "percent_agreement": float((both.nunique(axis=1) == 1).mean()) if len(both) else float("nan"),
        "fleiss_kappa": fleiss_kappa(both),
        "krippendorff_alpha": krippendorff_alpha_nominal(wide),
    }
    pairs = []
    for i in range(len(names)):
        for j in range(i + 1, len(names)):
            pair = wide[[names[i], names[j]]].dropna()
            if len(pair) == 0:
                continue
            labels = [s for s in (STATES if column == "gold_state" else sorted(set(pair.to_numpy().ravel())))]
            labels = labels + sorted(set(pair.to_numpy().ravel()) - set(labels))
            pairs.append(
                {
                    "a": names[i],
                    "b": names[j],
                    "n": int(len(pair)),
                    "percent_agreement": float((pair[names[i]] == pair[names[j]]).mean()),
                    "cohen_kappa": cohen_kappa(pair[names[i]].tolist(), pair[names[j]].tolist(), labels=labels),
                    "labels": labels,
                    "confusion": confusion_matrix(pair[names[i]].tolist(), pair[names[j]].tolist(), labels).tolist(),
                }
            )
    rep["pairs"] = pairs
    disagreements = both[both.nunique(axis=1) > 1]
    rep["disagreements"] = disagreements.reset_index().to_dict(orient="records")
    return rep


def agreement_to_markdown(rep: dict[str, Any], title: str = "Inter-annotator agreement") -> str:
    def f(x):
        return "n/a" if x is None or x != x else f"{x:.3f}"

    lines = [
        f"# {title}",
        "",
        f"Annotators: {', '.join(rep['annotators'])}; items: {rep['n_items']} "
        f"({rep['n_items_all_annotators']} rated by all). Column: `{rep['column']}`.",
        "",
        "| Measure | Value |",
        "|---|---:|",
        f"| Percent agreement (items rated by all) | {f(rep['percent_agreement'])} |",
        f"| Fleiss' kappa | {f(rep['fleiss_kappa'])} |",
        f"| Krippendorff's alpha (nominal) | {f(rep['krippendorff_alpha'])} |",
    ]
    for p in rep["pairs"]:
        lines += [
            "",
            f"## {p['a']} vs {p['b']} (n = {p['n']})",
            "",
            f"Percent agreement {f(p['percent_agreement'])}, Cohen's kappa {f(p['cohen_kappa'])}.",
            "",
            "| " + p["a"] + " \\ " + p["b"] + " | " + " | ".join(p["labels"]) + " |",
            "|---|" + "---:|" * len(p["labels"]),
        ]
        for lab, row in zip(p["labels"], p["confusion"]):
            lines.append(f"| {lab} | " + " | ".join(str(int(x)) for x in row) + " |")
    lines += ["", f"{len(rep['disagreements'])} items with disagreement (listed in the JSON report)."]
    return "\n".join(lines) + "\n"
