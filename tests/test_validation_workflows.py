"""Validation, indicators, quality checks and installation diagnostics."""

import contextlib
import io
import json
import tempfile
import unittest
from importlib import resources
from pathlib import Path

import numpy as np
import pandas as pd
from helpers import FakeResponse, FakeSession

from aidisclose import LocalSource, Pipeline, Results
from aidisclose.aggregate import firm_year_panel, group_panel, read_industry_map
from aidisclose.cli import main
from aidisclose.pipeline import DOCUMENT_COLUMNS, SENTENCE_COLUMNS
from aidisclose.validate import (
    LABEL_STUDIO_CONFIG,
    agreement_report,
    estimate_recall,
    fleiss_kappa,
    from_label_studio,
    krippendorff_alpha_nominal,
    read_annotation_file,
    sample_recall,
    to_label_studio,
    write_annotation_file,
)
from aidisclose.validate.external import eurostat_ai_share_by_industry, load_benchmark


def run(*args) -> str:
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        main([str(a) for a in args])
    return buf.getvalue()


def demo_results(**kw) -> Results:
    d = Path(str(resources.files("aidisclose.demo").joinpath("data")))
    src = LocalSource(d, metadata=d / "metadata.csv")
    return Pipeline(**kw).run(src.list_filings(), source=src)


def doc(fid, firm, year, sub, words=10000):
    row = {c: 0 if c.startswith("n_") else "" for c in DOCUMENT_COLUMNS}
    row.update(
        filing_id=fid,
        firm_id=firm,
        firm_name=firm,
        fiscal_year=year,
        country="US",
        status="ok",
        n_words=words,
        n_substantive=sub,
        n_keyword_sentences=max(sub, 1),
        n_baseline_sentences=1,
        source="t",
        extra="{}",
        industry_code="7372",
        industry_scheme="SIC",
    )
    return row


def sent(sid, fid, state, text, cats="general", reason="DEPLOYMENT"):
    row = {c: "" for c in SENTENCE_COLUMNS}
    row.update(
        sentence_id=sid,
        filing_id=fid,
        state=state,
        classification_status="ok",
        text=text,
        categories=cats,
        reason=reason,
        raw_terms="ai_acronym",
        confidence=0.8,
    )
    return row


class IndicatorTests(unittest.TestCase):
    def setUp(self):
        docs = pd.DataFrame(
            [
                doc("a21", "a", 2021, 0),
                doc("a22", "a", 2022, 2),
                doc("a23", "a", 2023, 2),
                doc("b22", "b", 2022, 1),
                doc("c21", "c", 2021, 0),
                doc("c23", "c", 2023, 1),
            ]
        )
        sents = pd.DataFrame(
            [
                sent("1", "a21", "WEAK", "We plan to use AI in our warehouses.", reason="FORWARD_LOOKING"),
                sent("2", "a22", "SUBSTANTIVE", "Our AI routing engine cut fuel use by 7% in 2022.", "general;ml"),
                sent(
                    "3",
                    "a22",
                    "SUBSTANTIVE",
                    "We launched a generative AI assistant for 2,000 agents.",
                    "genai",
                    "PRODUCT",
                ),
                sent("4", "a23", "SUBSTANTIVE", "Our AI routing engine cut fuel use by 9% in 2023.", "general;ml"),
                sent("5", "a23", "SUBSTANTIVE", "We acquired an AI start-up for $40 million.", "general", "INVESTMENT"),
                sent("6", "b22", "SUBSTANTIVE", "We use computer vision to inspect parcels.", "cv"),
                sent("7", "c23", "SUBSTANTIVE", "We use machine learning to price policies.", "ml"),
            ]
        )
        self.res = Results(sentences=sents, documents=docs)

    def test_breakdowns(self):
        p = firm_year_panel(self.res).set_index("filing_id")
        self.assertEqual(p.loc["a22", "n_substantive_genai"], 1)
        self.assertEqual(p.loc["a22", "n_substantive_ml"], 1)
        self.assertEqual(p.loc["a22", "n_substantive_product"], 1)
        self.assertEqual(p.loc["a23", "n_substantive_investment"], 1)
        self.assertEqual(p.loc["b22", "substantive_cv_any"], 1)
        self.assertEqual(p.loc["a21", "n_weak_forward_looking"], 1)

    def test_dynamics(self):
        p = firm_year_panel(self.res).set_index("filing_id")
        # restated routing claim with new figures repeats; the acquisition is new
        self.assertEqual(int(p.loc["a23", "n_substantive_new"]), 1)
        self.assertAlmostEqual(p.loc["a23", "substantive_novelty"], 0.5)
        self.assertEqual(int(p.loc["a22", "n_substantive_new"]), 2)
        self.assertTrue(pd.isna(p.loc["b22", "n_substantive_new"]))  # no previous year observed
        self.assertEqual(int(p.loc["a22", "substantive_onset"]), 1)
        self.assertEqual(int(p.loc["a23", "substantive_onset"]), 0)
        self.assertTrue(pd.isna(p.loc["b22", "substantive_onset"]))  # censored
        self.assertTrue(pd.isna(p.loc["c23", "substantive_onset"]))  # 2022 missing: gap
        self.assertEqual(int(p.loc["a21", "substantive_onset"]), 0)
        self.assertEqual(int(p.loc["a23", "first_substantive_year"]), 2022)
        self.assertEqual(int(p.loc["a23", "first_observed_year"]), 2021)

    def test_onset_only_in_first_substantive_year(self):
        docs = pd.DataFrame([doc("x20", "x", 2020, 1), doc("x21", "x", 2021, 0), doc("x22", "x", 2022, 1)])
        sents = pd.DataFrame(
            [
                sent("1", "x20", "SUBSTANTIVE", "We use AI to route trucks."),
                sent("2", "x22", "SUBSTANTIVE", "We use AI to price loans."),
            ]
        )
        p = firm_year_panel(Results(sentences=sents, documents=docs)).set_index("filing_id")
        self.assertTrue(pd.isna(p.loc["x20", "substantive_onset"]))  # censored at panel start
        self.assertEqual(int(p.loc["x22", "substantive_onset"]), 0)  # a return, not an onset

    def test_attributes_weights_and_map(self):
        attrs = pd.DataFrame({"firm_id": ["a", "b", "c"], "employees": [1000, 0, 3000], "country": ["X", "Y", "Z"]})
        p = firm_year_panel(self.res, firm_attributes=attrs)
        self.assertIn("attr_country", p.columns)  # clashing names are prefixed
        g = group_panel(p, by=["fiscal_year"], weight="employees")
        y23 = g[g["fiscal_year"] == 2023].iloc[0]
        self.assertEqual(int(y23["n_firms_weighted"]), 2)
        self.assertIn("share_substantive_genai", g.columns)
        with self.assertRaises(ValueError):
            group_panel(p, by=["fiscal_year"], weight="nope")
        with self.assertRaises(ValueError):
            firm_year_panel(self.res, firm_attributes=pd.DataFrame({"x": [1]}))
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "m.csv").write_text("industry_code,group\n7372,Software\n", encoding="utf-8")
            m = read_industry_map(Path(tmp, "m.csv"))
            self.assertEqual(firm_year_panel(self.res, industry_map=m)["nace_section"].unique().tolist(), ["Software"])


class RecallEstimatorTests(unittest.TestCase):
    def synthetic_pool(self):
        rng = np.random.default_rng(1)
        rows = []
        for d in range(40):
            for pool, size, drawn, rate in (("probe", 50, 10, 0.2), ("random", 2000, 5, 0.01)):
                for j in range(drawn):
                    truth = "SUBSTANTIVE" if rng.random() < rate else "NO_CLAIM"
                    rows.append(
                        {
                            "sentence_id": f"{d}-{pool}-{j}",
                            "filing_id": f"f{d}",
                            "pool": pool,
                            "stratum_size": size,
                            "n_drawn": drawn,
                            "inclusion_prob": drawn / size,
                            "text": "x",
                            "truth": truth,
                        }
                    )
        return pd.DataFrame(rows)

    def test_unbiased_on_full_annotation(self):
        pool = self.synthetic_pool()
        to_annotate, key = sample_recall(pool, n=len(pool), share_probe=0.5, seed=3)
        gold = key.copy()
        gold["gold_state"] = gold["truth"]
        rep = estimate_recall(gold, found_substantive=500, n_boot=200)
        est = rep["estimates"]
        # Horvitz-Thompson total equals the weighted count implied by the design
        expected = sum((1 / r.inclusion_prob) for r in pool.itertuples() if r.truth == "SUBSTANTIVE")
        self.assertAlmostEqual(est["missed_substantive"], expected, places=6)
        self.assertAlmostEqual(est["recall_substantive"], 500 / (500 + expected), places=6)
        self.assertAlmostEqual(est["non_retrieved_sentences"], 40 * (50 + 2000), places=6)
        lo, hi = rep["ci"]["recall_substantive"]
        self.assertLess(lo, est["recall_substantive"])
        self.assertGreater(hi, est["recall_substantive"])
        self.assertNotIn("gold_state", to_annotate.columns.difference(["gold_state"]))

    def test_sample_allocation_and_errors(self):
        pool = self.synthetic_pool()
        to_annotate, key = sample_recall(pool, n=100, share_probe=0.9, seed=0)
        self.assertEqual(len(to_annotate), 100)
        self.assertEqual((key["recall_stratum"] == "probe").sum(), 90)
        self.assertNotIn("pool", to_annotate.columns)
        with self.assertRaises(ValueError):
            sample_recall(pool.iloc[:0])
        with self.assertRaises(ValueError):
            estimate_recall(key.assign(gold_state="NO_CLAIM"))  # no found count
        main = pd.DataFrame({"gold_state": ["SUBSTANTIVE", "WEAK"], "sampling_weight": [10, 5]})
        rep = estimate_recall(key.assign(gold_state="NO_CLAIM"), main_gold=main, n_boot=50)
        self.assertEqual(rep["estimates"]["recall_substantive"], 1.0)
        self.assertEqual(rep["estimates"]["found_claim"], 15)


class AgreementTests(unittest.TestCase):
    def test_fleiss_wikipedia_example(self):
        counts = [
            [0, 0, 0, 0, 14],
            [0, 2, 6, 4, 2],
            [0, 0, 3, 5, 6],
            [0, 3, 9, 2, 0],
            [2, 2, 8, 1, 1],
            [7, 7, 0, 0, 0],
            [3, 2, 6, 3, 0],
            [2, 5, 3, 2, 2],
            [6, 5, 2, 1, 0],
            [0, 2, 2, 3, 7],
        ]
        rows = [[c for c, k in zip("ABCDE", r) for _ in range(k)] for r in counts]
        self.assertAlmostEqual(fleiss_kappa(pd.DataFrame(rows)), 0.210, places=3)

    def test_krippendorff_reference_example(self):
        nan = np.nan
        data = {
            "A": [1, 2, 3, 3, 2, 1, 4, 1, 2, nan, nan, nan],
            "B": [1, 2, 3, 3, 2, 2, 4, 1, 2, 5, nan, 3],
            "C": [nan, 3, 3, 3, 2, 3, 4, 2, 2, 5, 1, nan],
            "D": [1, 2, 3, 3, 2, 4, 4, 1, 2, 5, 1, nan],
        }
        self.assertAlmostEqual(krippendorff_alpha_nominal(pd.DataFrame(data)), 0.743, places=3)

    def test_report(self):
        a = pd.DataFrame(
            {
                "sentence_id": list("abcdef"),
                "gold_state": ["SUBSTANTIVE", "WEAK", "NO_CLAIM", "WEAK", "NO_CLAIM", "weak"],
            }
        )
        b = pd.DataFrame(
            {"sentence_id": list("abcdeg"), "gold_state": ["SUBSTANTIVE", "WEAK", "WEAK", "WEAK", "NO_CLAIM", "WEAK"]}
        )
        rep = agreement_report([a, b])
        self.assertEqual(rep["n_items_all_annotators"], 5)
        self.assertAlmostEqual(rep["percent_agreement"], 0.8)
        self.assertEqual(len(rep["pairs"]), 1)
        self.assertEqual(len(rep["disagreements"]), 1)
        self.assertTrue(-1 <= rep["pairs"][0]["cohen_kappa"] <= 1)
        with self.assertRaises(ValueError):
            agreement_report([a])


class AnnotationFileTests(unittest.TestCase):
    def sample(self):
        return pd.DataFrame(
            {
                "sentence_id": ["s1", "s2"],
                "firm_name": ["Ünïcode AG", "B"],
                "fiscal_year": [2023, 2024],
                "section": ["AR:risk", "10-K:1"],
                "prev_text": ["", "p"],
                "text": ["We use KI–AI.", "t"],
                "next_text": ["", ""],
                "gold_state": ["", ""],
                "gold_reason": ["", ""],
                "annotator": ["", ""],
                "notes": ["", ""],
            }
        )

    def test_xlsx_roundtrip_with_dropdowns(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as tmp:
            path = write_annotation_file(self.sample(), Path(tmp, "s.xlsx"))
            wb = load_workbook(path)
            ws = wb["annotate"]
            formulas = sorted(dv.formula1 for dv in ws.data_validations.dataValidation)
            self.assertEqual(len(formulas), 2)
            self.assertTrue(all(f.startswith("=lists!") for f in formulas))
            self.assertIn("codebook", wb.sheetnames)
            back = read_annotation_file(path)
            self.assertEqual(back["firm_name"].tolist(), ["Ünïcode AG", "B"])
            self.assertEqual(back["text"].iloc[0], "We use KI–AI.")

    def test_csv_has_bom(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = write_annotation_file(self.sample(), Path(tmp, "s.csv"))
            self.assertTrue(path.read_bytes().startswith(b"\xef\xbb\xbf"))
            self.assertEqual(read_annotation_file(path).columns[0], "sentence_id")

    def test_label_studio_roundtrip(self):
        tasks = to_label_studio(self.sample())
        self.assertEqual(tasks[0]["data"]["sentence_id"], "s1")
        self.assertIn('<Choices name="state" toName="target"', LABEL_STUDIO_CONFIG)
        export = [
            {
                "id": 1,
                "data": tasks[0]["data"],
                "annotations": [
                    {
                        "completed_by": 7,
                        "was_cancelled": False,
                        "result": [
                            {
                                "from_name": "state",
                                "to_name": "target",
                                "type": "choices",
                                "value": {"choices": ["WEAK"]},
                            },
                            {
                                "from_name": "reason",
                                "to_name": "target",
                                "type": "choices",
                                "value": {"choices": ["VAGUE"]},
                            },
                            {
                                "from_name": "notes",
                                "to_name": "target",
                                "type": "textarea",
                                "value": {"text": ["unsure"]},
                            },
                        ],
                    },
                    {
                        "completed_by": {"id": 8, "email": "b@x.org"},
                        "result": [
                            {
                                "from_name": "state",
                                "to_name": "target",
                                "type": "choices",
                                "value": {"choices": ["WEAK"]},
                            }
                        ],
                    },
                    {"completed_by": 9, "was_cancelled": True, "result": []},
                ],
            },
        ]
        df = from_label_studio(export)
        self.assertEqual(df["annotator"].tolist(), ["7", "b@x.org"])
        self.assertEqual(df.iloc[0][["gold_state", "gold_reason", "notes"]].tolist(), ["WEAK", "VAGUE", "unsure"])
        self.assertEqual(from_label_studio(json.dumps(export)).shape, df.shape)


class BenchmarkTests(unittest.TestCase):
    def test_eurostat_by_industry(self):
        doc = {
            "version": "2.0",
            "class": "dataset",
            "id": ["size_emp", "nace_r2", "geo", "time"],
            "size": [1, 3, 2, 1],
            "dimension": {
                "size_emp": {"category": {"index": {"GE10": 0}}},
                "nace_r2": {"category": {"index": {"C": 0, "C10-C12": 1, "J": 2}}},
                "geo": {"category": {"index": {"DK": 0, "EL": 1}}},
                "time": {"category": {"index": {"2023": 0}}},
            },
            "value": {"0": 10.0, "1": 8.0, "2": 12.0, "3": 9.0, "4": 40.0, "5": 30.0},
        }
        doc["id"] = ["indic_is", "unit"] + doc["id"]
        doc["size"] = [1, 1] + doc["size"]
        doc["dimension"].update(
            {"indic_is": {"category": {"index": {"E_AI_TANY": 0}}}, "unit": {"category": {"index": {"PC_ENT": 0}}}}
        )
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "e.json").write_text(json.dumps(doc), encoding="utf-8")
            df = eurostat_ai_share_by_industry(path=Path(tmp, "e.json"))
        self.assertEqual(sorted(df["nace_section"].unique()), ["C", "J"])
        self.assertEqual(sorted(df["country"].unique()), ["DK", "GR"])
        self.assertAlmostEqual(df[(df["country"] == "GR") & (df["nace_section"] == "J")]["eurostat_share"].iloc[0], 0.3)

    def test_generic_loader(self):
        with tempfile.TemporaryDirectory() as tmp:
            Path(tmp, "b.csv").write_text("geo,ref_year,ai_rate\nus,2023,5.4%\nde,2023,12\n", encoding="utf-8")
            df = load_benchmark(Path(tmp, "b.csv"), value="ai_rate", rename={"geo": "country", "ref_year": "year"})
            self.assertEqual(df["country"].tolist(), ["US", "DE"])
            self.assertAlmostEqual(df["benchmark_share"].iloc[0], 0.054)
            with self.assertRaises(ValueError):
                load_benchmark(Path(tmp, "b.csv"), value="nope")


class QADoctorTests(unittest.TestCase):
    def test_qa_on_demo(self):
        from aidisclose.qa import quality_report, quality_to_markdown

        rep = quality_report(demo_results())
        checks = {f["check"] for f in rep["findings"]}
        self.assertIn("skipped_language", checks)
        self.assertIn("no_recall_pool", checks)
        self.assertIn("Quality report", quality_to_markdown(rep))

    def test_qa_flags_problems(self):
        from aidisclose.qa import quality_report

        docs = pd.DataFrame(
            [
                doc("x1", "x", 2021, 0, words=0),
                doc("x2", "x", 2023, 0),
                {**doc("x3", "x", 2023, 0, words=500), "status": "error", "error": "HTTPError: 503"},
            ]
        )
        docs.loc[1, "form"] = "10-K"
        sents = pd.DataFrame([sent("x2#full#00001", "x2", "WEAK", "A" * 1200)])
        sents["section"] = "full"
        rep = quality_report(Results(sentences=sents, documents=docs))
        checks = {f["check"]: f for f in rep["findings"]}
        for name in ("failed_documents", "empty_text", "items_not_found", "long_sentences", "coverage_gaps"):
            self.assertIn(name, checks)
        self.assertGreaterEqual(rep["summary"]["n_warnings"], 3)

    def test_doctor(self):
        from aidisclose.doctor import check_environment, check_network, format_rows

        with tempfile.TemporaryDirectory() as tmp:
            rows = check_environment(Path(tmp) / "cache")
        names = {r["check"]: r["status"] for r in rows}
        self.assertEqual(names["pandas"], "ok")
        self.assertEqual(names["cache directory"], "ok")
        sess = FakeSession({"https://data.sec.gov/": b"{}", "https://filings.xbrl.org/": (403, b"no")})
        net = check_network("UA a@b.org", llm_base_url="http://localhost:9/v1", session=sess)
        status = {r["check"]: r["status"] for r in net}
        self.assertEqual(status["SEC EDGAR (submissions API)"], "ok")
        self.assertEqual(status["filings.xbrl.org"], "fail")
        self.assertTrue(any("localhost:9/v1/models" in u for u, _ in sess.calls))
        self.assertIn("pandas", format_rows(rows))
        self.assertIsInstance(FakeResponse(b"x").content, bytes)


class CLIWorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.out = Path(cls.tmp.name)
        d = Path(str(resources.files("aidisclose.demo").joinpath("data")))
        cls.res_dir = cls.out / "res"
        run(
            "run",
            "--source",
            "local",
            "--dir",
            d,
            "--metadata",
            d / "metadata.csv",
            "--out",
            cls.res_dir,
            "-q",
            "--recall-random",
            "3",
            "--recall-probe-cap",
            "3",
        )

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def test_sample_xlsx_and_label_studio(self):
        text = run(
            "sample",
            "--results",
            self.res_dir,
            "--n",
            "20",
            "--min-per-stratum",
            "3",
            "--out",
            self.out / "s.xlsx",
            "--label-studio",
            self.out / "ls.json",
        )
        self.assertIn("Label Studio tasks", text)
        self.assertTrue((self.out / "ls_config.xml").exists())
        self.assertEqual(
            len(json.loads((self.out / "ls.json").read_text())), len(read_annotation_file(self.out / "s.xlsx"))
        )

    def test_recall_workflow(self):
        run("sample-recall", "--results", self.res_dir, "--n", "15", "--out", self.out / "r.csv")
        sheet = read_annotation_file(self.out / "r.csv")
        sheet["gold_state"] = ["SUBSTANTIVE"] + ["NO_CLAIM"] * (len(sheet) - 1)
        sheet.to_csv(self.out / "r_done.csv", index=False)
        text = run(
            "evaluate-recall",
            "--gold",
            self.out / "r_done.csv",
            "--key",
            self.out / "r_key.csv",
            "--results",
            self.res_dir,
            "--n-boot",
            "50",
            "--out",
            self.out / "recall.md",
        )
        self.assertIn("Recall of SUBSTANTIVE sentences", text)
        rep = json.loads((self.out / "recall.json").read_text())
        self.assertGreater(rep["estimates"]["missed_substantive"], 0)
        self.assertEqual(rep["found_source"], "unvalidated classifier predictions")

    def test_agreement_and_label_studio_import(self):
        a = pd.DataFrame({"sentence_id": ["1", "2", "3"], "gold_state": ["WEAK", "WEAK", "NO_CLAIM"]})
        b = pd.DataFrame({"sentence_id": ["1", "2", "3"], "gold_state": ["WEAK", "SUBSTANTIVE", "NO_CLAIM"]})
        a.to_csv(self.out / "ann_alice.csv", index=False)
        b.to_csv(self.out / "ann_bob.csv", index=False)
        text = run(
            "agreement",
            "--annotations",
            self.out / "ann_alice.csv",
            self.out / "ann_bob.csv",
            "--out",
            self.out / "agree.md",
        )
        self.assertIn("ann_alice vs ann_bob", text)
        export = [
            {
                "data": {"sentence_id": "1"},
                "annotations": [
                    {"completed_by": 1, "result": [{"from_name": "state", "value": {"choices": ["WEAK"]}}]},
                    {"completed_by": 2, "result": [{"from_name": "state", "value": {"choices": ["NO_CLAIM"]}}]},
                ],
            }
        ]
        (self.out / "export.json").write_text(json.dumps(export))
        self.assertIn(
            "2 annotations by 2 annotators",
            run("import-label-studio", "--export", self.out / "export.json", "--out", self.out / "ls.csv"),
        )
        self.assertIn("Krippendorff", run("agreement", "--label-studio", self.out / "export.json"))

    def test_panel_options_qa_doctor_relabel(self):
        (self.out / "attrs.csv").write_text("firm_id,employees\nnorthwind,1200\nharbor,300\n", encoding="utf-8")
        run(
            "panel",
            "--results",
            self.res_dir,
            "--level",
            "year",
            "--firm-attributes",
            self.out / "attrs.csv",
            "--weight",
            "employees",
            "--out",
            self.out / "year.csv",
            "-q",
        )
        year = pd.read_csv(self.out / "year.csv")
        self.assertIn("n_firms_weighted", year.columns)
        text = run("qa", "--results", self.res_dir, "--out", self.out / "qa.md")
        self.assertIn("Quality report", text)
        self.assertTrue((self.out / "qa.json").exists())
        self.assertIn("pandas", run("doctor", "--cache-dir", self.out / "c"))

    def test_benchmark_command(self):
        run("panel", "--results", self.res_dir, "--level", "country-year", "--out", self.out / "cy.csv", "-q")
        cy = pd.read_csv(self.out / "cy.csv")
        bench = cy[["country", "fiscal_year"]].rename(columns={"fiscal_year": "yr"}).assign(rate=50)
        bench.to_csv(self.out / "bench.csv", index=False)
        text = run(
            "benchmark",
            "--panel",
            self.out / "cy.csv",
            "--file",
            self.out / "bench.csv",
            "--value",
            "rate",
            "--rename",
            "yr=year",
            "--min-firms",
            "1",
            "--out",
            self.out / "bm.csv",
        )
        self.assertIn("share_substantive", text)
        self.assertAlmostEqual(pd.read_csv(self.out / "bm.csv")["benchmark_share"].iloc[0], 0.5)


if __name__ == "__main__":
    unittest.main()
